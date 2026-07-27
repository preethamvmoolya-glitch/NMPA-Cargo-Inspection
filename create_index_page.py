import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def set_table_borders(table, color="000000", sz="4", val="single"):
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'<w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:left w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:right w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:insideV w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)

def make_row_header(row):
    trPr = row._tr.get_or_add_trPr()
    tblHeader = parse_xml(f'<w:tblHeader {nsdecls("w")}/>')
    cantSplit = parse_xml(f'<w:cantSplit {nsdecls("w")}/>')
    trPr.append(tblHeader)
    trPr.append(cantSplit)

def make_row_cant_split(row):
    trPr = row._tr.get_or_add_trPr()
    cantSplit = parse_xml(f'<w:cantSplit {nsdecls("w")}/>')
    trPr.append(cantSplit)

# Project-specific index items based on NMPA Cargo Inspection & AI RMS System
items = [
    ("Abstract", "Unnumbered"),
    ("Chapter 1: Synopsis & Introduction", "1"),
    ("1.1 Introduction of the System", "1"),
    ("1.1.1 Project Title", "1"),
    ("1.1.2 Category", "2"),
    ("1.1.3 Overview", "3"),
    ("1.1.4 Comparable Port Systems Analysis", "4"),
    ("1.2 Background", "5"),
    ("1.2.1 Introduction of New Mangalore Port Authority (NMPA)", "5"),
    ("1.2.2 Port Infrastructure & Cargo Operational Load", "6"),
    ("1.2.3 Agile Development Approach", "7"),
    ("1.3 Objectives of the System", "8"),
    ("1.4 Scope of the System", "9"),
    ("1.5 Structure of the System", "10"),
    ("1.6 System Architecture Overview (Figure 1.1)", "11"),
    ("1.7 End User Roles & Responsibilities", "12"),
    ("1.8 Software/Hardware Used for Development (Table 1.1)", "13"),
    ("1.9 Software/Hardware Required for Implementation (Table 1.2)", "14"),
    
    ("Chapter 2: Software Requirements Specification (SRS)", "15"),
    ("2.1 Introduction", "15"),
    ("2.1.1 Purpose", "15"),
    ("2.1.2 Scope", "16"),
    ("2.1.3 Definitions, Acronyms and Abbreviations", "17"),
    ("2.1.4 Regulatory Frameworks & Maritime Compliance References", "18"),
    ("2.2 Overall Description", "19"),
    ("2.2.1 Product Perspective", "19"),
    ("2.2.2 Product Functions", "20"),
    ("2.2.3 User Characteristics", "21"),
    ("2.2.4 General Constraints", "22"),
    ("2.2.5 Assumptions & Dependencies", "23"),
    ("2.3 Special Requirements (Bilingual UI, Dark/Light Mode, Local Storage Sync)", "24"),
    ("2.4 Functional Requirements (Table 2.1)", "25"),
    ("2.4.1 User Authentication & Role Administration", "25"),
    ("2.4.2 Import General Manifest (IGM) & Cargo Logging", "26"),
    ("2.4.3 AI Risk Management System (RMS) Threat Classification Engine", "27"),
    ("2.4.4 Weighbridge Physical Check & Tolerance Verification", "28"),
    ("2.4.5 Port Authority Clearance Adjudication & QR Gate Pass Generation", "29"),
    ("2.4.6 Public QR Verification Registry & Grievance Portal", "30"),
    ("2.5 Design Constraints & Data Standards", "31"),
    ("2.6 System Attributes (Reliability, Scalability, Security)", "32"),
    ("2.7 Other Requirements (Table 2.2)", "33"),
    
    ("Chapter 3: System Design & Architecture", "34"),
    ("3.1 Architectural Decomposition & Layered Design (Figure 3.1)", "34"),
    ("3.2 System Assumptions and Constraints", "35"),
    ("3.3 Functional Decomposition Breakdown", "36"),
    ("3.4 Description of Programs & Process Flow", "37"),
    ("3.4.1 Context Flow Diagram CFD (Figure 3.2)", "37"),
    ("3.4.2 Level 0 DFD - Overall System Data Flow (Figure 3.3)", "38"),
    ("3.4.3 Level 1 DFD - AI RMS Threat Evaluation Pipeline (Figure 3.4)", "39"),
    ("3.4.4 Level 2 DFD - Weighbridge Check & Authority Adjudication (Figure 3.5)", "40"),
    ("3.4.5 State Machine Diagram - Cargo Clearance Pipeline (Figure 3.6)", "41"),
    
    ("Chapter 4: Database Design & Entity Schemes", "42"),
    ("4.1 Introduction & Database Selection (SQLite / MongoDB)", "42"),
    ("4.2 Purpose and Scope", "43"),
    ("4.3 Table Definitions & Schemas (Tables 4.1, 4.2, 4.3, 4.4)", "44"),
    ("4.3.1 Users Table Schema (users)", "44"),
    ("4.3.2 Cargo Manifests Table Schema (manifests)", "45"),
    ("4.3.3 Weighbridge Physical Checks Table Schema (inspections)", "46"),
    ("4.3.4 Public Grievances Table Schema (grievances)", "47"),
    ("4.3.5 Security Audit Trail Table Schema (audit_logs)", "48"),
    ("4.4 Entity-Relationship (ER) Diagram (Figure 4.1)", "49"),
    
    ("Chapter 5: Detailed Design & Program Logic", "50"),
    ("5.1 Structure of the Software Package", "50"),
    ("5.2 Core Module Implementation", "51"),
    ("5.2.1 Authentication & Session Security (JWT, Bcrypt)", "51"),
    ("5.2.2 AI Risk Management System (RMS) Keyword Classifier", "52"),
    ("5.2.3 Weighbridge Scale Verification & Discrepancy Tolerance Engine", "53"),
    ("5.2.4 Port Authority Clearance Adjudication Console", "54"),
    ("5.2.5 Cryptographic QR Gate Pass Generator & Verification Portal", "55"),
    ("5.2.6 Public Grievance Redressal Portal & Chairman's Inbox", "56"),
    ("5.3 User Interface Gallery & Console Screenshots (Figures 5.1 - 5.6)", "57"),
    ("5.3.1 Frontline Inspector IGM Registration Screen (Figure 5.1)", "57"),
    ("5.3.2 AI RMS Threat Tier Notification Panel (Figure 5.2)", "58"),
    ("5.3.3 Weighbridge Inspection & Weight Variance Screen (Figure 5.3)", "59"),
    ("5.3.4 Port Authority Clearance Adjudication Dashboard (Figure 5.4)", "60"),
    ("5.3.5 QR Gate Pass Public Verification Portal (Figure 5.5)", "61"),
    ("5.3.6 Grievance Submission Portal & Chairman Office Queue (Figure 5.6)", "62"),
    
    ("Chapter 6: User Operations Manual & Guidelines", "63"),
    ("6.1 System Administrator Operations Manual", "63"),
    ("6.2 Frontline Port Inspector Operations Manual", "64"),
    ("6.3 Senior Port Authority Adjudication Manual", "65"),
    ("6.4 Public User & Grievance Portal Manual", "66"),
    
    ("Chapter 7: Testing and Quality Assurance", "67"),
    ("7.1 Testing Methodology & Strategy", "67"),
    ("7.1.1 Unit Testing (API Endpoints & RMS Logic)", "67"),
    ("7.1.2 Integration Testing (Weighbridge & Authority Flow)", "68"),
    ("7.1.3 System Workflow & Security Testing", "69"),
    ("7.2 Comprehensive Test Reports & Matrix (Tables 7.1, 7.2, 7.3, 7.4)", "70"),
    
    ("Chapter 8: Results, Discussion & Performance Benchmarks (Table 8.1)", "74"),
    ("Chapter 9: Conclusion & Scope for Future Enhancements", "77"),
    ("9.1 Conclusion", "77"),
    ("9.2 Operational Limitations", "78"),
    ("9.3 Scope for Future Enhancements", "79"),
    
    ("Abbreviations and Acronyms", "81"),
    ("Appendix A: NMPA Port Operations & Compliance Guidelines", "82"),
    ("Appendix B: Complete System Source Code Listings (Backend & Frontend)", "84"),
    ("Bibliography", "120")
]

def generate_index_docx(filename):
    doc = docx.Document()
    
    # Standard 1 inch Margins
    section = doc.sections[0]
    section.top_margin = Inches(1.0)
    section.bottom_margin = Inches(1.0)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)
    
    # Title "INDEX PAGE" - Bold, uppercase, pure black font, left aligned matching simple box design
    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    title_p.paragraph_format.space_before = Pt(0)
    title_p.paragraph_format.space_after = Pt(18)
    
    run_title = title_p.add_run("INDEX PAGE")
    run_title.font.name = "Arial"
    run_title.font.size = Pt(18)
    run_title.bold = True
    run_title.font.color.rgb = RGBColor(0, 0, 0) # Pure Black

    # Create Simple Grid Box Table
    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    
    # Set full simple black border box (top, bottom, left, right, insideH, insideV)
    set_table_borders(table, color="000000", sz="6", val="single")
    
    col_widths = [Inches(5.2), Inches(1.3)]
    
    # Header Row
    hdr_row = table.rows[0]
    make_row_header(hdr_row)
    
    cell_0 = hdr_row.cells[0]
    cell_1 = hdr_row.cells[1]
    
    cell_0.width = col_widths[0]
    cell_1.width = col_widths[1]
    
    set_cell_margins(cell_0, top=120, bottom=120, left=150, right=150)
    set_cell_margins(cell_1, top=120, bottom=120, left=150, right=150)
    
    p0 = cell_0.paragraphs[0]
    p0.paragraph_format.space_before = Pt(0)
    p0.paragraph_format.space_after = Pt(0)
    r0 = p0.add_run("System Topic / Figure / Table Index")
    r0.font.name = "Times New Roman"
    r0.font.size = Pt(12)
    r0.bold = True
    r0.font.color.rgb = RGBColor(0, 0, 0) # Pure Black
    
    p1 = cell_1.paragraphs[0]
    p1.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p1.paragraph_format.space_before = Pt(0)
    p1.paragraph_format.space_after = Pt(0)
    r1 = p1.add_run("Page No.")
    r1.font.name = "Times New Roman"
    r1.font.size = Pt(12)
    r1.bold = True
    r1.font.color.rgb = RGBColor(0, 0, 0) # Pure Black

    # Populate Data Rows
    for topic, pg in items:
        row = table.add_row()
        make_row_cant_split(row)
        
        c0 = row.cells[0]
        c1 = row.cells[1]
        
        c0.width = col_widths[0]
        c1.width = col_widths[1]
        
        set_cell_margins(c0, top=100, bottom=100, left=150, right=150)
        set_cell_margins(c1, top=100, bottom=100, left=150, right=150)
            
        p_c0 = c0.paragraphs[0]
        p_c0.paragraph_format.space_before = Pt(0)
        p_c0.paragraph_format.space_after = Pt(0)
        p_c0.paragraph_format.line_spacing = 1.15
        
        r_c0 = p_c0.add_run(topic)
        r_c0.font.name = "Times New Roman"
        r_c0.font.size = Pt(12)
        r_c0.font.color.rgb = RGBColor(0, 0, 0) # Pure Black

        p_c1 = c1.paragraphs[0]
        p_c1.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p_c1.paragraph_format.space_before = Pt(0)
        p_c1.paragraph_format.space_after = Pt(0)
        p_c1.paragraph_format.line_spacing = 1.15
        
        r_c1 = p_c1.add_run(pg)
        r_c1.font.name = "Times New Roman"
        r_c1.font.size = Pt(12)
        r_c1.font.color.rgb = RGBColor(0, 0, 0) # Pure Black

    try:
        doc.save(filename)
        print(f"Document saved successfully as {filename}")
    except PermissionError:
        alt_name = "NMPA_Index_Page_Project.docx"
        doc.save(alt_name)
        print(f"File {filename} was locked. Saved as {alt_name} instead.")

if __name__ == "__main__":
    generate_index_docx("NMPA_Index_Page_Project.docx")
