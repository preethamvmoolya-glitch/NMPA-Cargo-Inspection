import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

def generate_abstract_docx(filename):
    doc = docx.Document()
    
    # 1-inch Standard Margins for University Academic Report
    section = doc.sections[0]
    section.top_margin = Inches(1.0)
    section.bottom_margin = Inches(1.0)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)
    
    # Heading: ABSTRACT
    heading_p = doc.add_paragraph()
    heading_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    heading_p.paragraph_format.space_before = Pt(0)
    heading_p.paragraph_format.space_after = Pt(24)
    
    run_heading = heading_p.add_run("ABSTRACT")
    run_heading.font.name = "Times New Roman"
    run_heading.font.size = Pt(18)
    run_heading.bold = True
    run_heading.font.color.rgb = RGBColor(0, 0, 0)
    
    # Paragraphs tailored specifically to the NMPA Cargo Inspection & AI RMS Project
    paragraphs = [
        "The maritime port sector serves as the vital gateway for international commerce and logistics, where rapid, secure, and compliant cargo clearance is essential to maximize port operational throughput and prevent anchorage congestion. At the New Mangalore Port Authority (NMPA) situated in Panambur, Mangaluru, managing substantial annual vessel and cargo volume requires transitioning away from traditional, labor-intensive paper-reliant inspection workflows. Manual processing, unverified container weight declarations, and subjective risk evaluations introduce vulnerabilities such as revenue leakage, security threats, weight fraud, and prolonged vessel turnaround times.",
        
        "To resolve these operational challenges, this academic project report presents the complete design, software architecture, implementation, and operational deployment of the New Mangalore Port Authority (NMPA) Cargo Inspection and AI Risk Management System. The platform establishes a centralized digital governance portal that automates Import General Manifest (IGM) registration, real-time weighbridge physical verification, artificial intelligence threat profiling, senior port authority adjudication, and tamper-evident QR Code gate pass generation.",
        
        "The core innovation of the system lies in its integrated AI Risk Management System (RMS), driven by a deterministic Natural Language Processing (NLP) classification engine. Upon manifest entry by frontline port inspectors, the AI engine dynamically categorizes incoming cargo descriptions into three threat tiers: Tier 1 Critical Risk (Red Tag) for volatile energy fuels, chemicals, and bulk minerals requiring mandatory regulatory auditing; Tier 2 Elevated Risk (Yellow Tag) for perishable agricultural commodities such as cashews and coffee or high-tariff electronics requiring phytosanitary and tax validation; and Tier 3 Routine Risk (Green Tag) for standard general cargo. Additionally, the platform automates physical weighbridge validation by cross-referencing declared container tonnage against actual scale measurements, triggering automated warnings whenever weight discrepancies exceed administrator-defined tolerance thresholds.",
        
        "Architecturally built using a modern full-stack web architecture comprising a React.js single-page frontend (built with Vite) and a Node.js/Express.js RESTful API backend integrated with SQLite storage (cargo.db), the system enforces strict Role-Based Access Control (RBAC) across System Administrators, Port Inspectors, and Senior Port Authorities. Upon clearance adjudication by the Port Authority, the system issues a cryptographically verifiable digital QR Gate Pass for public security gate verification. Furthermore, a dedicated Public Grievance Redressal portal enables port operators and stakeholders to submit operational feedback directly to the NMPA Chairman's Office.",
        
        "Systematic workflow verification and end-to-end testing demonstrate that the NMPA Cargo Inspection System successfully eliminates paperwork bottlenecks, standardizes cargo threat profiling, mitigates weight discrepancy risks, and drastically reduces cargo clearance processing times, establishing a modernized digital standard for port authority operations."
    ]
    
    for p_text in paragraphs:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(12)
        p.paragraph_format.line_spacing = 1.5
        p.paragraph_format.first_line_indent = Inches(0.5)
        
        run = p.add_run(p_text)
        run.font.name = "Times New Roman"
        run.font.size = Pt(12)
        run.font.color.rgb = RGBColor(0, 0, 0)
        
    # Keywords Paragraph
    kw_p = doc.add_paragraph()
    kw_p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    kw_p.paragraph_format.space_before = Pt(18)
    kw_p.paragraph_format.space_after = Pt(0)
    kw_p.paragraph_format.line_spacing = 1.5
    
    r_kw_title = kw_p.add_run("Keywords: ")
    r_kw_title.font.name = "Times New Roman"
    r_kw_title.font.size = Pt(12)
    r_kw_title.bold = True
    r_kw_title.font.color.rgb = RGBColor(0, 0, 0)
    
    r_kw_text = kw_p.add_run("Cargo Inspection System, AI Risk Management System (RMS), Natural Language Processing (NLP), New Mangalore Port Authority (NMPA), Weighbridge Verification, QR Gate Pass, Role-Based Access Control (RBAC), Full-Stack Web Development, React.js, Node.js, Express.js, SQLite.")
    r_kw_text.font.name = "Times New Roman"
    r_kw_text.font.size = Pt(12)
    r_kw_text.italic = True
    r_kw_text.font.color.rgb = RGBColor(0, 0, 0)

    # Try saving to primary file or fallback if locked
    saved_path = filename
    try:
        doc.save(filename)
        print(f"Abstract Word document saved successfully as {filename}")
    except PermissionError:
        saved_path = "NMPA_Abstract_Page_Final.docx"
        doc.save(saved_path)
        print(f"File {filename} was locked. Saved as {saved_path} instead.")

if __name__ == "__main__":
    generate_abstract_docx("NMPA_Abstract_Page.docx")
